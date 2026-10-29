"""Exportación a Excel y plantilla oficial de carga."""
import io

from ..config import FIELDS, LABELS

def xlsx_bytes(rows, title='Datos'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title=title[:31]
    headers=FIELDS if not rows else [f for f in FIELDS if f in rows[0]]
    ws.append([LABELS.get(h,h) for h in headers])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='173B68'); c.alignment=Alignment(horizontal='center')
    for r in rows: ws.append([r.get(h,'') for h in headers])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for i,h in enumerate(headers,1):
        vals=[ws.cell(j,i).value for j in range(1,ws.max_row+1)]
        ws.column_dimensions[get_column_letter(i)].width=min(44,max(12,max(len(str(v or '')) for v in vals)+2))
    buf=io.BytesIO(); wb.save(buf); return buf.getvalue()

def template_bytes():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title='Plantilla de carga'
    ws.append([LABELS[f] for f in FIELDS])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='173B68'); c.alignment=Alignment(horizontal='center')
    ws.append([2026,'','','','COLOMBIA',1,'','','','','','','','','Presencial','Pregrado',1])
    for i,f in enumerate(FIELDS,1): ws.column_dimensions[get_column_letter(i)].width=min(34,max(14,len(LABELS[f])+3))
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(FIELDS))}2'
    notes=wb.create_sheet('Instrucciones')
    notes.append(['Campo','Qué debes colocar','Obligatorio'])
    for c in notes[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='173B68')
    required={'Ano','Nombre Programa','Sexo','Sede','Facultad','Modalidad','Tipo Programa'}
    for f in FIELDS: notes.append([LABELS[f], 'Usa el mismo formato de los datos existentes. Deja vacío si no aplica.', 'Sí' if f in required else 'No'])
    notes.append([]); notes.append(['Importante','No cambies los nombres de las columnas de la hoja Plantilla de carga. La aplicación las usa para integrar los nuevos registros.',''])
    for col in range(1,4): notes.column_dimensions[get_column_letter(col)].width=[28,72,16][col-1]
    return_bytes=io.BytesIO(); wb.save(return_bytes); return return_bytes.getvalue()

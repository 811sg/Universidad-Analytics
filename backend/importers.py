"""Importación de archivos (CSV, JSON, XLSX) y validación de filas."""
import csv
import io
import json
from pathlib import Path

from .config import LABELS, REQUIRED_IMPORT_FIELDS
from .schema import HEADER_ALIASES, key_norm, clean_value, normalize_row, is_valid

def read_multipart(handler):
    from email.parser import BytesParser
    from email.policy import default
    length=int(handler.headers.get("Content-Length","0")); body=handler.rfile.read(length)
    content_type=handler.headers.get("Content-Type","")
    if not content_type.lower().startswith("multipart/form-data"):
        raise ValueError("La importación requiere multipart/form-data.")
    msg=BytesParser(policy=default).parsebytes(("Content-Type: "+content_type+"\r\nMIME-Version: 1.0\r\n\r\n").encode()+body)
    item=next((p for p in msg.iter_parts() if p.get_param("name",header="content-disposition")=="file"),None)
    if item is None: raise ValueError("No se recibió ningún archivo.")
    return item.get_filename() or "", item.get_payload(decode=True) or b""

def canonicalize_import_row(row):
    if not isinstance(row, dict):
        return None
    out = {}
    for raw_key, value in row.items():
        field = HEADER_ALIASES.get(key_norm(raw_key))
        if field:
            out[field] = value
    # A completely blank spreadsheet row is ignored.
    if not any(clean_value(v) is not None for v in out.values()):
        return None
    return normalize_row(out)

def validate_import_rows(rows):
    accepted, rejected = [], []
    for number, row in enumerate(rows, start=2):
        if not isinstance(row, dict):
            rejected.append({"fila": number, "motivo": "Registro no válido."})
            continue
        canonical = canonicalize_import_row(row)
        if canonical is None:
            continue
        missing = [LABELS[f] for f in REQUIRED_IMPORT_FIELDS if not is_valid(canonical.get(f))]
        if missing:
            rejected.append({"fila": number, "motivo": "Faltan campos obligatorios: " + ", ".join(missing)})
            continue
        accepted.append(canonical)
    if not accepted and rows:
        detail = rejected[0]["motivo"] if rejected else "No se encontraron registros válidos."
        raise ValueError("El archivo no contiene registros válidos. " + detail)
    return accepted, rejected

def parse_upload(filename,raw):
    ext=Path(filename).suffix.lower()
    if ext=='.json':
        rows=json.loads(raw.decode('utf-8-sig'))
        if isinstance(rows,dict) and 'data' in rows: rows=rows['data']
    elif ext=='.csv':
        rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        # Prefer the official template sheet, otherwise use the first sheet.
        sheet='Plantilla de carga' if 'Plantilla de carga' in wb.sheetnames else wb.sheetnames[0]
        ws=wb[sheet]; vals=list(ws.values)
        if not vals: return [], []
        headers=[str(x).strip() if x is not None else '' for x in vals[0]]
        rows=[dict(zip(headers,row)) for row in vals[1:]]
    else: raise ValueError('Formato no compatible. Usa CSV, JSON o XLSX.')
    if not isinstance(rows,list): raise ValueError('El archivo debe contener una lista de registros.')
    return validate_import_rows(rows)
